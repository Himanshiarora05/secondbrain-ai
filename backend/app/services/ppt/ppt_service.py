"""
PowerPoint (.pptx) Text Extraction Service.

Extracts slide titles, body text, tables, and speaker notes with slide-number markers.
"""

from pathlib import Path
from pptx import Presentation


class PPTService:

    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from a PowerPoint (.pptx) presentation.

        Iterates through slides in order, extracting slide titles, shape text,
        table contents, and speaker notes. Formats with clear [Slide X] markers.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise ValueError(f"Failed to parse PowerPoint presentation: {e}") from e

        slides_text = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_parts = []
            title_text = ""

            # Check if slide has a title shape
            if slide.shapes.title and slide.shapes.title.text.strip():
                title_text = slide.shapes.title.text.strip()
                slide_parts.append(f"[Slide {slide_num}: {title_text}]")
            else:
                slide_parts.append(f"[Slide {slide_num}]")

            # Extract text from all other shapes and tables
            body_parts = []
            for shape in slide.shapes:
                # Skip title shape since already captured
                if shape == slide.shapes.title:
                    continue

                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        body_parts.append(text)
                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            table_rows.append(" | ".join(row_text))
                    if table_rows:
                        body_parts.append("\n".join(table_rows))

            if body_parts:
                slide_parts.append("\n".join(body_parts))

            # Extract speaker notes if present
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes = notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_parts.append(f"[Notes: {notes}]")

            # Only append slide if it contained some text
            joined_slide = "\n".join(slide_parts).strip()
            if len(slide_parts) > 1 or title_text:
                slides_text.append(joined_slide)

        return "\n\n".join(slides_text)
