"""Offline checks for old Office formats (.ppt/.doc) in the upload endpoints (no DB, no Chroma data, no network).

Run from backend/:  .venv/Scripts/python.exe tests/test_office_formats.py
"""
import asyncio
import io
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch, MagicMock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# Never touch the real database, vector store or model hub.
os.environ.setdefault("DATABASE_URL", "postgresql://offline:offline@127.0.0.1:1/offline")
os.environ["CHROMA_DIR"] = tempfile.mkdtemp(prefix="sb-test-chroma-")
os.environ["ANONYMIZED_TELEMETRY"] = "False"
os.environ["HF_HUB_OFFLINE"] = "1"

from docx import Document as WordDocument
from fastapi import HTTPException, UploadFile
from pptx import Presentation

from types import SimpleNamespace

# The signed-in user the route functions are called for (routes take it from get_current_user).
TEST_USER = SimpleNamespace(id=1)

from app.api import upload

OLE_BYTES = upload.OLE_SIGNATURE + b"\x00" * 504  # start of an OLE compound file


def _pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Cell Biology"
    slide.placeholders[1].text = "Mitochondria produce ATP through oxidative phosphorylation."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    doc = WordDocument()
    doc.add_heading("Cell Biology", level=1)
    doc.add_paragraph("Mitochondria produce ATP through oxidative phosphorylation.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _call(endpoint, filename, data):
    """Run an upload endpoint with a temp uploads dir and storage mocked out.

    Returns (result or HTTPException, files written, store mock).
    """
    tmp = Path(tempfile.mkdtemp(prefix="sb-test-uploads-"))
    file = UploadFile(file=io.BytesIO(data), filename=filename)
    with patch.object(upload, "UPLOAD_DIR", tmp), \
         patch.object(upload, "_store_document_and_chunks", return_value={"status": "stored"}) as store:
        try:
            result = asyncio.run(endpoint(file=file, db=MagicMock(), user=TEST_USER))
        except HTTPException as e:
            result = e
    return result, list(tmp.iterdir()), store


def _assert_rejected(result, written, store, message):
    assert isinstance(result, HTTPException), f"expected a rejection, got {result!r}"
    assert result.status_code == 400
    assert result.detail == message, result.detail
    assert written == [], f"rejected upload left files behind: {written}"
    store.assert_not_called()


def test_ppt_extension_rejected_before_saving():
    for name in ("lecture.ppt", "LECTURE.PPT"):
        _assert_rejected(*_call(upload.upload_pptx, name, OLE_BYTES), upload.LEGACY_PPT_MESSAGE)
    assert "Save As" in upload.LEGACY_PPT_MESSAGE and ".pptx" in upload.LEGACY_PPT_MESSAGE


def test_doc_extension_rejected_before_saving():
    for name in ("notes.doc", "NOTES.DOC"):
        _assert_rejected(*_call(upload.upload_docx, name, OLE_BYTES), upload.LEGACY_DOC_MESSAGE)
    assert "Save As" in upload.LEGACY_DOC_MESSAGE and ".docx" in upload.LEGACY_DOC_MESSAGE


def test_old_format_renamed_to_new_extension_is_explained():
    _assert_rejected(*_call(upload.upload_pptx, "renamed.pptx", OLE_BYTES), upload.OLE_PPTX_MESSAGE)
    _assert_rejected(*_call(upload.upload_docx, "renamed.docx", OLE_BYTES), upload.OLE_DOCX_MESSAGE)


def test_other_extensions_still_get_the_generic_message():
    result, written, _ = _call(upload.upload_pptx, "slides.key", b"x")
    assert result.detail == "Only PowerPoint (.pptx) files are allowed" and written == []
    result, written, _ = _call(upload.upload_docx, "notes.odt", b"x")
    assert result.detail == "Only Word (.docx) files are allowed" and written == []


def test_real_pptx_still_uploads():
    result, written, store = _call(upload.upload_pptx, "lecture.pptx", _pptx_bytes())
    assert result == {"status": "stored"}, result
    assert len(written) == 1 and written[0].suffix == ".pptx"
    kwargs = store.call_args.kwargs
    assert kwargs["source_type"] == "pptx" and "oxidative phosphorylation" in kwargs["content"]


def test_real_docx_still_uploads():
    result, written, store = _call(upload.upload_docx, "notes.docx", _docx_bytes())
    assert result == {"status": "stored"}, result
    assert len(written) == 1 and written[0].suffix == ".docx"
    kwargs = store.call_args.kwargs
    assert kwargs["source_type"] == "docx" and "oxidative phosphorylation" in kwargs["content"]


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"PASS {t.__name__}")
    print(f"\n{len(tests)} passed")
