"""Offline checks that a failed upload doesn't leave its file in uploads/, and that scanned
PDFs get their own message (temp uploads dir, storage mocked; no DB, Chroma data or network).

Run from backend/:  .venv/Scripts/python.exe tests/test_upload_cleanup.py
"""
import asyncio
import io
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# Never touch the real database, vector store or model hub.
os.environ.setdefault("DATABASE_URL", "postgresql://offline:offline@127.0.0.1:1/offline")
os.environ["CHROMA_DIR"] = tempfile.mkdtemp(prefix="sb-test-chroma-")
os.environ["ANONYMIZED_TELEMETRY"] = "False"
os.environ["HF_HUB_OFFLINE"] = "1"

import fitz
from docx import Document as WordDocument
from fastapi import HTTPException, UploadFile
from pptx import Presentation

from types import SimpleNamespace

# The signed-in user the route functions are called for (routes take it from get_current_user).
TEST_USER = SimpleNamespace(id=1)

from app.api import upload

STORED = {"status": "stored"}


def pdf_bytes(text="Mitochondria make ATP by oxidative phosphorylation. " * 5, image=False):
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_textbox(fitz.Rect(40, 40, 560, 800), text, fontsize=10)
    if image:
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
        pix.clear_with(200)
        page.insert_image(fitz.Rect(40, 40, 400, 400), stream=pix.tobytes("png"))
    data = doc.tobytes()
    doc.close()
    return data


def pptx_bytes(text="Mitochondria produce ATP."):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Cell Biology"
    slide.placeholders[1].text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def docx_bytes(text="Mitochondria produce ATP."):
    doc = WordDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def call(endpoint, filename, data, store=STORED, patches=()):
    """Run an upload endpoint against a temp uploads dir. Returns (result or HTTPException, files left)."""
    tmp = Path(tempfile.mkdtemp(prefix="sb-test-uploads-"))
    file = UploadFile(file=io.BytesIO(data), filename=filename)
    store_mock = MagicMock(side_effect=store) if isinstance(store, BaseException) else MagicMock(return_value=store)
    with patch.object(upload, "UPLOAD_DIR", tmp), patch.object(upload, "_store_document_and_chunks", store_mock), \
         patch.object(upload, "logger"):
        for target, attr, value in patches:
            patch.object(target, attr, value).start()
        try:
            result = asyncio.run(endpoint(file=file, db=MagicMock(), user=TEST_USER))
        except HTTPException as e:
            result = e
        finally:
            patch.stopall()
    return result, [p.name for p in tmp.iterdir()]


def rejected(result, status, detail=None):
    assert isinstance(result, HTTPException), f"expected an error, got {result!r}"
    assert result.status_code == status, (result.status_code, result.detail)
    if detail is not None:
        assert result.detail == detail, result.detail


def test_successful_uploads_keep_their_file():
    for endpoint, name, data, ext in [(upload.upload_pdf, "a.pdf", pdf_bytes(), ".pdf"),
                                      (upload.upload_pptx, "a.pptx", pptx_bytes(), ".pptx"),
                                      (upload.upload_docx, "a.docx", docx_bytes(), ".docx")]:
        result, left = call(endpoint, name, data)
        assert result == STORED, result
        assert len(left) == 1 and left[0].endswith(ext), left


def test_scanned_pdf_is_explained_and_removed():
    result, left = call(upload.upload_pdf, "scan.pdf", pdf_bytes(text="", image=True))
    rejected(result, 400, upload.SCANNED_PDF_MESSAGE)
    assert "scanned" in result.detail and "OCR" in result.detail
    assert left == [], left


def test_blank_pdf_keeps_the_plain_message_and_is_removed():
    result, left = call(upload.upload_pdf, "blank.pdf", pdf_bytes(text=""))
    rejected(result, 400, "No readable text found in this PDF file")
    assert left == [], left


def test_unreadable_pdf_is_removed():
    result, left = call(upload.upload_pdf, "broken.pdf", b"%PDF-1.4 this is not really a pdf")
    rejected(result, 400, "Could not extract text from this PDF file, please try another file")
    assert left == [], left


def test_no_chunks_is_removed():
    for endpoint, name, data, target in [(upload.upload_pdf, "a.pdf", pdf_bytes(), "chunk_pages"),
                                         (upload.upload_pptx, "a.pptx", pptx_bytes(), "chunk_text"),
                                         (upload.upload_docx, "a.docx", docx_bytes(), "chunk_text")]:
        result, left = call(endpoint, name, data, patches=[(upload.RAGService, target, staticmethod(lambda *a, **k: []))])
        rejected(result, 400, "No usable text chunks after processing")
        assert left == [], (name, left)


def test_storage_or_embedding_failure_is_removed():
    for endpoint, name, data in [(upload.upload_pdf, "a.pdf", pdf_bytes()),
                                 (upload.upload_pptx, "a.pptx", pptx_bytes()),
                                 (upload.upload_docx, "a.docx", docx_bytes())]:
        for status in (502, 500):
            result, left = call(endpoint, name, data, store=HTTPException(status_code=status, detail="boom"))
            rejected(result, status)
            assert left == [], (name, status, left)


def test_unexpected_crash_is_removed_and_still_raised():
    result, left = None, None
    try:
        result, left = call(upload.upload_pdf, "a.pdf", pdf_bytes(), store=RuntimeError("crash"))
        raise AssertionError("the crash should propagate")
    except RuntimeError:
        pass
    # Check directly: a crash inside the guarded block removes the file and re-raises.
    tmp = Path(tempfile.mkdtemp(prefix="sb-test-uploads-")) / "x.pdf"
    tmp.write_bytes(b"x")
    try:
        with upload._removed_on_failure(tmp):
            raise RuntimeError("crash")
    except RuntimeError:
        pass
    assert not tmp.exists()


def test_bad_office_files_are_removed():
    for endpoint, name in [(upload.upload_pptx, "broken.pptx"), (upload.upload_docx, "broken.docx")]:
        result, left = call(endpoint, name, b"PK\x03\x04 not a real zip")
        rejected(result, 400)
        assert left == [], (name, left)
    result, left = call(upload.upload_pptx, "empty.pptx", pptx_bytes(text=" "),
                        patches=[(upload.PPTService, "extract_text", staticmethod(lambda p: "   "))])
    rejected(result, 400, "No readable text found in this PowerPoint presentation")
    assert left == []
    result, left = call(upload.upload_docx, "odd.docx", docx_bytes(),
                        patches=[(upload.DOCXService, "extract_text", staticmethod(lambda p: (_ for _ in ()).throw(OSError("disk"))))])
    rejected(result, 502)
    assert left == []


def test_a_file_that_cannot_be_removed_does_not_hide_the_error():
    with patch.object(Path, "unlink", side_effect=PermissionError("locked")):
        result, _ = call(upload.upload_pdf, "blank.pdf", pdf_bytes(text=""))
    rejected(result, 400, "No readable text found in this PDF file")


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"PASS {t.__name__}")
    print(f"\n{len(tests)} passed")
