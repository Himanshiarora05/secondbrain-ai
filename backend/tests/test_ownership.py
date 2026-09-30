"""Offline checks that every endpoint needs a session and only ever touches the signed-in
user's data (the real app via TestClient, in-memory SQLite, a throwaway Chroma; embeddings,
YouTube and the AI mocked; no network).

Run from backend/:  .venv/Scripts/python.exe tests/test_ownership.py
"""
import io
import os
import sys
import tempfile
import warnings
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# Never touch the real database, vector store or model hub. (Set before main is
# imported: load_dotenv doesn't override variables that are already set.)
os.environ["DATABASE_URL"] = "postgresql://offline:offline@127.0.0.1:1/offline"
os.environ["CHROMA_DIR"] = tempfile.mkdtemp(prefix="sb-test-chroma-")
os.environ["ANONYMIZED_TELEMETRY"] = "False"
os.environ["HF_HUB_OFFLINE"] = "1"
warnings.filterwarnings("ignore", category=DeprecationWarning)

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import main
from app.api import upload
from app.database.chroma import get_collection
from app.database.db import Base, get_db
from app.models import Document
from app.routes import search as search_route
from app.services import search_service

PASSWORD = "correct horse battery"
PUBLIC = {("POST", "/api/v1/auth/signup"), ("POST", "/api/v1/auth/login"), ("POST", "/api/v1/auth/logout"),
          ("GET", "/api/v1/health"), ("GET", "/"),
          # Forgot password: used while signed out by definition (tests/test_password_reset.py).
          ("POST", "/api/v1/auth/password-reset/request"), ("POST", "/api/v1/auth/password-reset/check"),
          ("POST", "/api/v1/auth/password-reset/confirm")}
VIDEO = "https://www.youtube.com/watch?v=abcdefghijk"


def fake_embeddings(texts):
    return [[1.0, float(len(t) % 7), 0.5, 0.25] for t in texts]


engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)


@event.listens_for(engine, "connect")
def _fk_on(conn, _):
    conn.execute("PRAGMA foreign_keys=ON")


Base.metadata.create_all(engine)
Local = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def _db():
    db = Local()
    try:
        yield db
    finally:
        db.close()


main.app.dependency_overrides[get_db] = _db
patch.object(upload, "get_embeddings", side_effect=fake_embeddings).start()
patch.object(search_service, "get_embedding", side_effect=lambda q: fake_embeddings([q])[0]).start()
patch.object(search_route, "generate_answer", return_value="An answer.").start()


def account(email):
    client = TestClient(main.app)
    r = client.post("/api/v1/auth/signup", json={"email": email, "password": PASSWORD})
    assert r.status_code == 201, r.text
    client.user_id = r.json()["user"]["id"]
    return client


def docx(text):
    doc = WordDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_doc(client, name, text):
    tmp = Path(tempfile.mkdtemp(prefix="sb-test-uploads-"))
    with patch.object(upload, "UPLOAD_DIR", tmp):
        r = client.post("/api/v1/upload/docx", files={"file": (name, docx(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, r.text
    return r.json()["document_id"]


alice = account("alice@example.com")
bob = account("bob@example.com")
a1 = upload_doc(alice, "alice-cells.docx", "Mitochondria make ATP by oxidative phosphorylation in the inner membrane.")
a2 = upload_doc(alice, "alice-genes.docx", "Genes are stretches of DNA that code for proteins through transcription.")
b1 = upload_doc(bob, "bob-history.docx", "The printing press spread books across Europe in the fifteenth century.")


def test_every_endpoint_needs_a_session():
    # Every endpoint the app publishes (its OpenAPI schema, the same list /docs
    # shows), so one added later without get_current_user fails here.
    anon = TestClient(main.app)
    checked = []
    for route_path, operations in main.app.openapi()["paths"].items():
        path = route_path.replace("{document_id}", str(a1)).replace("{set_id}", "1")
        for method in (m.upper() for m in operations):
            if (method, route_path) in PUBLIC:
                continue
            r = anon.request(method, path, json={}, params={"query": "x", "count": 1})
            assert r.status_code == 401, f"{method} {route_path} answered {r.status_code} without a session"
            checked.append(f"{method} {route_path}")
    assert len(checked) >= 22, checked
    print(f"    ({len(checked)} protected endpoints checked)")


def test_uploads_belong_to_the_uploader():
    db = Local()
    assert {d.id: d.user_id for d in db.query(Document)} == {a1: alice.user_id, a2: alice.user_id, b1: bob.user_id}
    metas = get_collection().get(where={"document_id": a1}, include=["metadatas"])["metadatas"]
    assert metas and all(m["user_id"] == alice.user_id for m in metas), metas


def test_each_user_lists_only_their_documents():
    assert sorted(d["id"] for d in alice.get("/api/v1/documents").json()) == sorted([a1, a2])
    assert [d["id"] for d in bob.get("/api/v1/documents").json()] == [b1]


def test_another_users_document_looks_missing():
    for method, suffix, body in [
        ("GET", "/summary", None),
        ("POST", "/summary", None),
        ("GET", "/flashcards", None),
        ("POST", "/flashcards", None),
        ("PATCH", "", {"filename": "stolen.docx"}),
        ("DELETE", "", None),
    ]:
        r = bob.request(method, f"/api/v1/documents/{a1}{suffix}", json=body)
        missing = bob.request(method, f"/api/v1/documents/999999{suffix}", json=body)
        assert r.status_code == 404 and r.json() == missing.json(), f"{method} {suffix}: {r.status_code} {r.text} vs {missing.text}"
    doc = Local().get(Document, a1)
    assert doc is not None and doc.filename == "alice-cells.docx", "Alice's document is untouched"


def test_merged_sets_are_private():
    r = alice.post("/api/v1/merged-sets", json={"document_ids": [a1, a2]})
    assert r.status_code == 200, r.text
    set_id = r.json()["merged_set"]["id"]
    assert bob.get("/api/v1/merged-sets").json() == []
    for method, path, body in [
        ("GET", f"/api/v1/merged-sets/{set_id}", None),
        ("PATCH", f"/api/v1/merged-sets/{set_id}", {"name": "mine now"}),
        ("DELETE", f"/api/v1/merged-sets/{set_id}", None),
        ("GET", f"/api/v1/merged-sets/{set_id}/summary", None),
        ("POST", f"/api/v1/merged-sets/{set_id}/summary", None),
        ("GET", f"/api/v1/merged-sets/{set_id}/flashcards", None),
        ("POST", f"/api/v1/merged-sets/{set_id}/flashcards", None),
    ]:
        r = bob.request(method, path, json=body)
        assert r.status_code == 404 and r.json()["detail"] == "Merged set not found", f"{method} {path}: {r.status_code}"
    # Bob can't build a set from Alice's documents, or learn they exist.
    r = bob.post("/api/v1/merged-sets", json={"document_ids": [b1, a1]})
    assert r.status_code == 404 and r.json()["detail"] == f"Document not found: {a1}"
    assert alice.get(f"/api/v1/merged-sets/{set_id}").json()["name"] == "alice-cells.docx + alice-genes.docx"


def test_search_only_sees_your_own_documents():
    for client, own in ((alice, {"alice-cells.docx", "alice-genes.docx"}), (bob, {"bob-history.docx"})):
        r = client.get("/api/v1/search", params={"query": "proteins and books"})
        assert r.status_code == 200, r.text
        found = {m["document"] for m in r.json()["top_matches"]}
        assert found and found <= own, (found, own)


def test_duplicate_links_are_per_user():
    segments = [{"text": "Alright, so here we are in front of the elephants.", "start": 1.2, "duration": 3.0}]
    with patch.object(upload.YouTubeService, "fetch_transcript", return_value=("abcdefghijk", segments)):
        assert alice.post("/api/v1/upload/youtube", json={"url": VIDEO}).status_code == 200
        assert bob.post("/api/v1/upload/youtube", json={"url": VIDEO}).status_code == 200, "Bob can save the same video"
        again = alice.post("/api/v1/upload/youtube", json={"url": VIDEO})
    assert again.status_code == 409 and "already in your library" in again.json()["detail"]["message"]


def test_documents_from_before_accounts_are_hidden_until_assigned():
    db = Local()
    db.add(Document(file_id="legacy", filename="legacy.pdf", content="old", source_type="pdf"))
    db.commit()
    legacy = db.query(Document).filter(Document.file_id == "legacy").one().id
    for client in (alice, bob):
        assert legacy not in [d["id"] for d in client.get("/api/v1/documents").json()]
        assert client.get(f"/api/v1/documents/{legacy}/summary").status_code == 404


def test_logging_out_locks_the_api_again():
    carol = account("carol@example.com")
    assert carol.get("/api/v1/documents").status_code == 200
    carol.post("/api/v1/auth/logout")
    assert carol.get("/api/v1/documents").status_code == 401


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"PASS {t.__name__}")
    print(f"\n{len(tests)} passed")
